from __future__ import annotations

import queue
import threading
import uuid
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from src import OUTPUT_DIR, iso_now

DOWNLOAD_LOG_PATH = OUTPUT_DIR / "download_log.xlsx"
LIVE_FRAMEWORK_PATH = OUTPUT_DIR / "RAG_Framework_Live.xlsx"
CONVERGENCE_REPORT_PATH = OUTPUT_DIR / "convergence_report.xlsx"

DOWNLOAD_LOG_HEADERS = [
    "Document_ID",
    "Company_Name",
    "URL",
    "File_Type",
    "File_Path",
    "File_Size_MB",
    "Downloaded (Yes/No)",
    "Download_Status",
    "Download_Date",
    "Final_URL",
    "Was_Redirected",
    "HTTP_Status",
    "Source_Backend",
    "Response_Time_MS",
    "Content_Fingerprint",
    "Is_Duplicate",
    "Duplicate_Of",
    "Relevance_Score",
    "Priority_Domain (Yes/No)",
    "DDG_Title",
    "DDG_Snippet",
    "Query_Family",
    "Query_Stage",
    "Source_Type",
    "Source_Priority",
    "Confidence_Score",
    "Rejection_Reason",
    "Run_ID",
    "Notes",
    "Storage_Backend",
    "Bucket_Name",
    "Object_Key",
    "Cloud_URI",
    "Upload_Status",
]

REGISTRY_HEADERS = [
    "Document_ID",
    "Document_Name",
    "Document_Type",
    "Category",
    "Sub_Category",
    "Source",
    "Source_URL",
    "Query_Used",
    "Query_Family",
    "Query_Stage",
    "File_Path",
    "File_Size_MB",
    "Date_Acquired",
    "Processing_Status",
    "Rejection_Reason",
    "Company_Name",
    "Content_Fingerprint",
    "ETag",
    "Content_Last_Modified",
    "Notes",
    "Storage_Backend",
    "Bucket_Name",
    "Object_Key",
    "Cloud_URI",
    "Upload_Status",
]

PROCESSING_LOG_HEADERS = [
    "Log_ID",
    "Document_ID",
    "Document_Name",
    "Stage_ID",
    "Stage_Name",
    "Status",
    "Start_Time",
    "End_Time",
    "Records_Processed",
    "Error_Message",
    "Notes",
]

CONVERGENCE_SUMMARY_HEADERS = [
    "Company_Name",
    "Total_Queries",
    "Converged_Queries",
    "Query_Set_Hash",
    "All_Converged",
    "Last_Updated",
]

RUN_HISTORY_HEADERS = [
    "Run_ID",
    "Run_Number",
    "Date",
    "Total_URLs",
    "New_URLs",
    "Downloads_OK",
    "Downloads_Failed",
    "Duplicates_Skipped",
    "Convergence_Pct",
    "Duration_Minutes",
    "Notes",
]


class ExcelWriterService:
    def __init__(self):
        self.queue: queue.Queue[dict[str, Any]] = queue.Queue()
        self.thread: threading.Thread | None = None
        self._started = False

    def start(self) -> None:
        if self._started:
            return
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        self._ensure_workbooks()
        self.thread = threading.Thread(target=self._worker, name="excel-writer", daemon=True)
        self.thread.start()
        self._started = True

    def stop(self) -> None:
        if not self._started:
            return
        self.queue.put({"op": "__stop__"})
        self.queue.join()
        if self.thread is not None:
            self.thread.join()
        self.thread = None
        self._started = False

    def enqueue(self, item: dict[str, Any]) -> None:
        if not self._started:
            raise RuntimeError("Writer thread has not been started.")
        self.queue.put(item)

    def _ensure_workbooks(self) -> None:
        self._ensure_workbook(DOWNLOAD_LOG_PATH, {"Download_Log": DOWNLOAD_LOG_HEADERS})
        self._ensure_workbook(
            LIVE_FRAMEWORK_PATH,
            {
                "Document_Registry": REGISTRY_HEADERS,
                "Processing_Log": PROCESSING_LOG_HEADERS,
                "Convergence_Summary": CONVERGENCE_SUMMARY_HEADERS,
                "Run_History": RUN_HISTORY_HEADERS,
            },
        )

    def _ensure_workbook(self, path: Path, sheets: dict[str, list[str]]) -> None:
        if path.exists():
            workbook = load_workbook(path)
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)
        existing = set(workbook.sheetnames)
        for sheet_name, headers in sheets.items():
            if sheet_name not in existing:
                sheet = workbook.create_sheet(sheet_name)
                sheet.append(headers)
        workbook.save(path)
        workbook.close()

    def _worker(self) -> None:
        while True:
            item = self.queue.get()
            try:
                op = item["op"]
                if op == "__stop__":
                    return
                if op == "download_row":
                    self._append_row(DOWNLOAD_LOG_PATH, "Download_Log", DOWNLOAD_LOG_HEADERS, item["row"])
                elif op == "registry_row":
                    self._append_row(LIVE_FRAMEWORK_PATH, "Document_Registry", REGISTRY_HEADERS, item["row"])
                elif op == "processing_log":
                    self._append_row(LIVE_FRAMEWORK_PATH, "Processing_Log", PROCESSING_LOG_HEADERS, item["row"])
                elif op == "convergence_update":
                    self._replace_rows(
                        LIVE_FRAMEWORK_PATH,
                        "Convergence_Summary",
                        CONVERGENCE_SUMMARY_HEADERS,
                        item["rows"],
                    )
                elif op == "run_summary":
                    self._append_row(LIVE_FRAMEWORK_PATH, "Run_History", RUN_HISTORY_HEADERS, item["row"])
                elif op == "convergence_report":
                    self._write_convergence_report(item["report"])
            finally:
                self.queue.task_done()

    def _append_row(self, path: Path, sheet_name: str, headers: list[str], row: dict[str, Any]) -> None:
        workbook = load_workbook(path)
        try:
            sheet = workbook[sheet_name]
            sheet.append([row.get(header) for header in headers])
            workbook.save(path)
        finally:
            workbook.close()

    def _replace_rows(self, path: Path, sheet_name: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
        workbook = load_workbook(path)
        try:
            sheet = workbook[sheet_name]
            if sheet.max_row > 1:
                sheet.delete_rows(2, sheet.max_row - 1)
            for row in rows:
                sheet.append([row.get(header) for header in headers])
            workbook.save(path)
        finally:
            workbook.close()

    def _write_convergence_report(self, report: dict[str, Any]) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)

        company_sheet = workbook.create_sheet("Company_Convergence")
        company_sheet.append(CONVERGENCE_SUMMARY_HEADERS)
        for row in report.get("company_rows", []):
            company_sheet.append([row.get(header) for header in CONVERGENCE_SUMMARY_HEADERS])

        low_yield_sheet = workbook.create_sheet("Low_Yield_Queries")
        low_yield_headers = [
            "company_name",
            "query_text",
            "query_family",
            "run_count",
            "total_urls_found",
            "total_new_urls_found",
            "yield_ratio",
        ]
        low_yield_sheet.append(low_yield_headers)
        for row in report.get("low_yield_rows", []):
            low_yield_sheet.append([row.get(header) for header in low_yield_headers])

        summary_sheet = workbook.create_sheet("Project_Summary")
        summary_sheet.append(["Metric", "Value"])
        for key, value in report.get("project_summary", {}).items():
            summary_sheet.append([key, value])

        workbook.save(CONVERGENCE_REPORT_PATH)
        workbook.close()


_SERVICE = ExcelWriterService()


def start_writer_thread() -> None:
    _SERVICE.start()


def stop_writer_thread() -> None:
    _SERVICE.stop()


def enqueue_download_row(result_dict: dict[str, Any]) -> None:
    row = {
        "Document_ID": result_dict.get("doc_id"),
        "Company_Name": result_dict.get("company_name"),
        "URL": result_dict.get("url"),
        "File_Type": result_dict.get("file_type"),
        "File_Path": result_dict.get("file_path"),
        "File_Size_MB": result_dict.get("file_size_mb"),
        "Downloaded (Yes/No)": result_dict.get("downloaded"),
        "Download_Status": result_dict.get("download_status"),
        "Download_Date": result_dict.get("downloaded_at", iso_now()),
        "Final_URL": result_dict.get("final_url"),
        "Was_Redirected": result_dict.get("was_redirected"),
        "HTTP_Status": result_dict.get("http_status"),
        "Source_Backend": result_dict.get("source_backend"),
        "Response_Time_MS": result_dict.get("response_time_ms"),
        "Content_Fingerprint": result_dict.get("content_fingerprint"),
        "Is_Duplicate": "Yes" if result_dict.get("is_duplicate") else "No",
        "Duplicate_Of": result_dict.get("duplicate_of_url"),
        "Relevance_Score": result_dict.get("relevance_score"),
        "Priority_Domain (Yes/No)": "Yes" if result_dict.get("priority_domain") else "No",
        "DDG_Title": result_dict.get("ddg_title"),
        "DDG_Snippet": result_dict.get("ddg_snippet"),
        "Query_Family": result_dict.get("query_family"),
        "Query_Stage": result_dict.get("query_stage"),
        "Source_Type": result_dict.get("source_type"),
        "Source_Priority": result_dict.get("source_priority"),
        "Confidence_Score": result_dict.get("confidence_score"),
        "Rejection_Reason": result_dict.get("rejection_reason"),
        "Run_ID": result_dict.get("run_id"),
        "Notes": result_dict.get("notes"),
        "Storage_Backend": result_dict.get("storage_backend"),
        "Bucket_Name": result_dict.get("bucket_name"),
        "Object_Key": result_dict.get("object_key"),
        "Cloud_URI": result_dict.get("cloud_uri"),
        "Upload_Status": result_dict.get("upload_status"),
    }
    _SERVICE.enqueue({"op": "download_row", "row": row})


def enqueue_registry_row(result_dict: dict[str, Any]) -> None:
    file_path = result_dict.get("file_path")
    document_name = Path(file_path).name if file_path else result_dict.get("doc_id")
    row = {
        "Document_ID": result_dict.get("doc_id"),
        "Document_Name": document_name,
        "Document_Type": result_dict.get("file_type"),
        "Category": result_dict.get("category"),
        "Sub_Category": result_dict.get("sub_category"),
        "Source": result_dict.get("source_name") or result_dict.get("source_backend"),
        "Source_URL": result_dict.get("url"),
        "Query_Used": result_dict.get("discovered_by_query") or result_dict.get("query_used"),
        "Query_Family": result_dict.get("query_family"),
        "Query_Stage": result_dict.get("query_stage"),
        "File_Path": file_path,
        "File_Size_MB": result_dict.get("file_size_mb"),
        "Date_Acquired": result_dict.get("downloaded_at", iso_now()),
        "Processing_Status": result_dict.get("download_status"),
        "Rejection_Reason": result_dict.get("rejection_reason"),
        "Company_Name": result_dict.get("company_name"),
        "Content_Fingerprint": result_dict.get("content_fingerprint"),
        "ETag": result_dict.get("etag"),
        "Content_Last_Modified": result_dict.get("content_last_modified"),
        "Notes": result_dict.get("notes"),
        "Storage_Backend": result_dict.get("storage_backend"),
        "Bucket_Name": result_dict.get("bucket_name"),
        "Object_Key": result_dict.get("object_key"),
        "Cloud_URI": result_dict.get("cloud_uri"),
        "Upload_Status": result_dict.get("upload_status"),
    }
    _SERVICE.enqueue({"op": "registry_row", "row": row})


def enqueue_processing_log(result_dict: dict[str, Any]) -> None:
    row = {
        "Log_ID": result_dict.get("log_id") or str(uuid.uuid4()),
        "Document_ID": result_dict.get("doc_id"),
        "Document_Name": result_dict.get("document_name") or result_dict.get("doc_id"),
        "Stage_ID": result_dict.get("stage_id"),
        "Stage_Name": result_dict.get("stage_name"),
        "Status": result_dict.get("status"),
        "Start_Time": result_dict.get("start_time", iso_now()),
        "End_Time": result_dict.get("end_time", iso_now()),
        "Records_Processed": result_dict.get("records_processed", 0),
        "Error_Message": result_dict.get("error_message"),
        "Notes": result_dict.get("notes"),
    }
    _SERVICE.enqueue({"op": "processing_log", "row": row})


def enqueue_convergence_update(stats: list[dict[str, Any]]) -> None:
    rows = []
    for item in stats:
        rows.append(
            {
                "Company_Name": item.get("company_name"),
                "Total_Queries": item.get("total_queries"),
                "Converged_Queries": item.get("converged_queries"),
                "Query_Set_Hash": item.get("query_set_hash"),
                "All_Converged": item.get("all_converged"),
                "Last_Updated": item.get("last_updated"),
            }
        )
    _SERVICE.enqueue({"op": "convergence_update", "rows": rows})


def enqueue_run_summary(run_stats: dict[str, Any]) -> None:
    row = {
        "Run_ID": run_stats.get("run_id"),
        "Run_Number": run_stats.get("run_number"),
        "Date": run_stats.get("completed_at") or run_stats.get("started_at"),
        "Total_URLs": run_stats.get("total_urls_seen"),
        "New_URLs": run_stats.get("new_urls_found"),
        "Downloads_OK": run_stats.get("downloaded_ok"),
        "Downloads_Failed": run_stats.get("failed"),
        "Duplicates_Skipped": run_stats.get("duplicates_skipped"),
        "Convergence_Pct": run_stats.get("project_convergence_pct"),
        "Duration_Minutes": run_stats.get("duration_minutes"),
        "Notes": run_stats.get("notes"),
    }
    _SERVICE.enqueue({"op": "run_summary", "row": row})


def enqueue_convergence_report(report: dict[str, Any]) -> None:
    _SERVICE.enqueue({"op": "convergence_report", "report": report})
